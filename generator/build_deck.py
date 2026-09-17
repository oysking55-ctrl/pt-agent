#!/usr/bin/env python3
"""Build a PPTX deck and a matching Reveal.js HTML deck from a content JSON file.

Usage:
    python build_deck.py content/01-visual-studio-agent.json

Outputs (relative to repo root):
    docs/<slug>/slides.pptx
    docs/<slug>/index.html
Also updates docs/index.html with a card linking to the new deck.
"""
import json
import re
import sys
import html
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

REPO_ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = REPO_ROOT / "docs"

NAVY = RGBColor(0x11, 0x1C, 0x3B)
ACCENT = RGBColor(0x4F, 0x8C, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x33, 0x3A, 0x4D)


def load_content(path: Path) -> dict:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# PPTX build
# ---------------------------------------------------------------------------

def set_background(slide, color: RGBColor):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_title_slide(prs: Presentation, topic: str, subtitle: str, audience: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)

    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = topic
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE

    box2 = slide.shapes.add_textbox(Inches(0.8), Inches(4.3), Inches(11.7), Inches(1))
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(24)
    p2.font.color.rgb = ACCENT

    if audience:
        box3 = slide.shapes.add_textbox(Inches(0.8), Inches(4.9), Inches(11.7), Inches(0.6))
        p3 = box3.text_frame.paragraphs[0]
        p3.text = f"대상: {audience}"
        p3.font.size = Pt(18)
        p3.font.color.rgb = RGBColor(0xCC, 0xD6, 0xF6)

    rule = slide.shapes.add_shape(1, Inches(0.85), Inches(2.15), Inches(1.4), Pt(4))
    rule.fill.solid()
    rule.fill.fore_color.rgb = ACCENT
    rule.line.fill.background()
    return slide


def add_terms_footer(slide, terms):
    if not terms:
        return
    line = slide.shapes.add_shape(1, Inches(0.7), Inches(6.55), Inches(11.9), Pt(1.2))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0xD8, 0xDC, 0xE6)
    line.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.7), Inches(6.65), Inches(11.9), Inches(0.75))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]

    label = p.add_run()
    label.text = "용어 설명   "
    label.font.size = Pt(12)
    label.font.bold = True
    label.font.color.rgb = ACCENT

    for i, term in enumerate(terms):
        sep = p.add_run()
        sep.text = "" if i == 0 else "     ·     "
        sep.font.size = Pt(12)
        sep.font.color.rgb = SLATE

        name = p.add_run()
        name.text = f"{term['term']} "
        name.font.size = Pt(12)
        name.font.bold = True
        name.font.color.rgb = NAVY

        definition = p.add_run()
        definition.text = term["def"]
        definition.font.size = Pt(12)
        definition.font.color.rgb = SLATE


def add_flow_diagram(slide, steps, top):
    """Draw a horizontal box-and-arrow flow diagram; returns the bottom edge (Emu)."""
    n = len(steps)
    total_w_in = 11.9
    arrow_w_in = 0.45
    height_in = 0.9
    box_w_in = (total_w_in - arrow_w_in * (n - 1)) / n
    left_in = 0.7

    for i, step in enumerate(steps):
        shape = slide.shapes.add_shape(5, Inches(left_in), top, Inches(box_w_in), Inches(height_in))
        shape.fill.solid()
        shape.fill.fore_color.rgb = ACCENT
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Pt(4)
        tf.margin_right = Pt(4)
        p = tf.paragraphs[0]
        p.text = step
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE
        left_in += box_w_in

        if i < n - 1:
            arrow_box = slide.shapes.add_textbox(Inches(left_in), top, Inches(arrow_w_in), Inches(height_in))
            atf = arrow_box.text_frame
            atf.vertical_anchor = MSO_ANCHOR.MIDDLE
            ap = atf.paragraphs[0]
            ap.text = "→"
            ap.alignment = PP_ALIGN.CENTER
            ap.font.size = Pt(20)
            ap.font.bold = True
            ap.font.color.rgb = NAVY
            left_in += arrow_w_in

    return top + Inches(height_in)


def add_content_slide(prs: Presentation, index: int, total: int, title: str, bullets, code: str = None, terms=None, diagram=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)

    # top accent bar
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(11.9), Inches(1.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = NAVY

    body_top = Inches(1.55)
    body_height = Inches(4.9) if terms else Inches(5.9)

    bullets_top = body_top
    if diagram:
        diagram_bottom = add_flow_diagram(slide, diagram, body_top)
        bullets_top = diagram_bottom + Inches(0.3)
        body_height = body_height - (bullets_top - body_top)

    if code:
        body_width = Inches(7.0)
    else:
        body_width = Inches(11.9)

    body_box = slide.shapes.add_textbox(Inches(0.7), bullets_top, body_width, body_height)
    tf = body_box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {bullet}"
        p.font.size = Pt(19)
        p.font.color.rgb = SLATE
        p.space_after = Pt(12)

    if code:
        code_box = slide.shapes.add_textbox(Inches(7.9), body_top, Inches(4.7), Inches(3.2))
        code_box.fill.solid()
        code_box.fill.fore_color.rgb = NAVY
        tf = code_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(12)
        tf.margin_top = Pt(12)
        for i, line in enumerate(code.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(14)
            p.font.name = "Consolas"
            p.font.color.rgb = RGBColor(0x9C, 0xD8, 0xFF)

    add_terms_footer(slide, terms)

    # page number
    pn = slide.shapes.add_textbox(Inches(12.5), Inches(7.05), Inches(0.6), Inches(0.35))
    p = pn.text_frame.paragraphs[0]
    p.text = f"{index}/{total}"
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE
    return slide


def build_pptx(data: dict, out_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs, data["topic"], data.get("subtitle", ""), data.get("audience", ""))

    sections = data["sections"]
    total = len(sections) + 2  # title + closing
    for i, sec in enumerate(sections, start=2):
        add_content_slide(prs, i, total, sec["title"], sec["bullets"], sec.get("code"), sec.get("terms"), sec.get("diagram"))

    closing = data.get("closing")
    if closing:
        add_content_slide(prs, total, total, closing["title"], closing["bullets"], terms=closing.get("terms"))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)


# ---------------------------------------------------------------------------
# Reveal.js HTML build
# ---------------------------------------------------------------------------

REVEAL_CDN = "https://cdnjs.cloudflare.com/ajax/libs/reveal.js/5.1.0"

HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{topic}</title>
<link rel="stylesheet" href="{cdn}/reveal.css">
<link rel="stylesheet" href="{cdn}/theme/night.css" id="theme">
<link rel="stylesheet" href="{cdn}/plugin/highlight/monokai.css">
<style>
  :root {{
    --accent: #4f8cff;
    --r-main-font-size: 34px;
    --r-heading1-size: 2em;
    --r-heading2-size: 1.35em;
    --r-block-margin: 14px;
  }}
  .reveal h1, .reveal h2 {{ color: var(--accent); }}
  .reveal .slides section {{
    text-align: left;
    padding-top: 64px;
    box-sizing: border-box;
    height: 100%;
    display: flex;
    flex-direction: column;
  }}
  .reveal .slides section > h2 {{ margin-top: 0; margin-bottom: 0.5em; flex: 0 0 auto; }}
  .reveal .slide-body {{ flex: 1 1 auto; min-height: 0; overflow-y: auto; }}
  .reveal ul {{ display: block; font-size: 0.76em; line-height: 1.4; margin: 0; padding-left: 0.9em; }}
  .reveal ul.dense {{ font-size: 0.64em; line-height: 1.3; }}
  .reveal ul.sparse {{ font-size: 0.88em; line-height: 1.5; }}
  .reveal li {{ margin-bottom: 0.5em; }}
  .reveal section pre {{ font-size: 0.46em; width: 100%; box-shadow: none; }}
  .reveal .flow {{ display: flex; align-items: stretch; gap: 0; margin: 0 0 20px; flex-wrap: wrap; flex: 0 0 auto; }}
  .reveal .flow-step {{
    flex: 1 1 0; min-width: 110px; display: flex; align-items: center; justify-content: center;
    background: var(--accent); color: #0b1020; font-weight: bold; font-size: 0.56em; line-height: 1.3;
    text-align: center; border-radius: 10px; padding: 14px 10px;
  }}
  .reveal .flow-arrow {{ flex: 0 0 auto; display: flex; align-items: center; padding: 0 10px; font-size: 0.75em; color: var(--accent); }}
  .reveal section aside.notes {{ display: none; }}
  .reveal .term-footer {{
    flex: 0 0 auto;
    margin-top: 10px;
    padding-top: 10px;
    border-top: 1px solid rgba(255,255,255,0.18);
    font-size: 0.4em;
    line-height: 1.5;
    color: #b8c2e0;
  }}
  .reveal .term-footer b {{ color: #e8ecf7; }}
  .reveal .term-footer .term-label {{ color: var(--accent); font-weight: bold; margin-right: 8px; }}
  .reveal .title-slide {{ text-align: center; padding-top: 0; justify-content: center; }}
  .reveal .title-slide h1 {{ font-size: 1.5em; }}
  .reveal .title-slide h3 {{ color: #ccd6f6; font-weight: 400; font-size: 0.9em; }}
  .reveal .title-slide .audience-tag {{ color: #9aa4c2; font-size: 0.55em; margin-top: 1em; }}
  .topbar {{
    position: fixed; top: 0; left: 0; right: 0; padding: 10px 20px;
    display: flex; justify-content: flex-end; gap: 10px; z-index: 100;
  }}
  .topbar a {{
    background: var(--accent); color: #fff; text-decoration: none;
    padding: 8px 16px; border-radius: 6px; font-size: 14px; font-family: sans-serif;
  }}
  .topbar a:hover {{ opacity: 0.85; }}
</style>
</head>
<body>
<div class="topbar">
  <a href="slides.pptx" download>PPT 다운로드</a>
  <a href="../">← 목록으로</a>
</div>
<div class="reveal">
  <div class="slides">
{slides}
  </div>
</div>
<script src="{cdn}/reveal.js"></script>
<script src="{cdn}/plugin/notes/notes.js"></script>
<script src="{cdn}/plugin/highlight/highlight.js"></script>
<script>
  Reveal.initialize({{
    width: 1280,
    height: 720,
    margin: 0.04,
    minScale: 0.2,
    maxScale: 1.5,
    center: false,
    hash: true,
    slideNumber: true,
    transition: 'slide',
    plugins: [ RevealNotes, RevealHighlight ]
  }});
</script>
</body>
</html>
"""


def esc(s: str) -> str:
    return html.escape(s, quote=False)


def density_class(bullets, has_code=False) -> str:
    """Pick a font-size class so long/many bullets still fit on one slide."""
    total_chars = sum(len(b) for b in bullets)
    budget = total_chars * (1.6 if has_code else 1.0)
    if budget > 220 or len(bullets) >= 6:
        return "dense"
    if budget < 120 and len(bullets) <= 4:
        return "sparse"
    return ""


def bullets_ul(bullets, extra_class="", style=""):
    cls = " ".join(c for c in [density_class(bullets, bool(style)), extra_class] if c)
    class_attr = f' class="{cls}"' if cls else ""
    style_attr = f' style="{style}"' if style else ""
    items = "\n".join(f"        <li>{esc(b)}</li>" for b in bullets)
    return f"      <ul{class_attr}{style_attr}>\n{items}\n      </ul>"


def flow_diagram_html(steps):
    parts = []
    for i, step in enumerate(steps):
        parts.append(f'<div class="flow-step">{esc(step)}</div>')
        if i < len(steps) - 1:
            parts.append('<div class="flow-arrow">→</div>')
    items = "\n".join(f"        {p}" for p in parts)
    return f"      <div class=\"flow\">\n{items}\n      </div>"


def terms_footer_html(terms):
    if not terms:
        return ""
    parts = [f'<b>{esc(t["term"])}</b> {esc(t["def"])}' for t in terms]
    body = "&nbsp;&nbsp;·&nbsp;&nbsp;".join(parts)
    return f'\n      <div class="term-footer"><span class="term-label">용어 설명</span>{body}</div>'


def build_html(data: dict, out_path: Path):
    slides_html = []

    audience = data.get("audience", "")
    audience_html = f'\n      <p class="audience-tag">대상: {esc(audience)}</p>' if audience else ""
    slides_html.append(f"""    <section class="title-slide">
      <h1>{esc(data['topic'])}</h1>
      <h3>{esc(data.get('subtitle', ''))}</h3>{audience_html}
    </section>""")

    for sec in data["sections"]:
        notes = sec.get("notes", "")
        notes_html = f"\n      <aside class=\"notes\">{esc(notes)}</aside>" if notes else ""
        code = sec.get("code")
        diagram = sec.get("diagram")
        if diagram:
            body = flow_diagram_html(diagram) + "\n" + bullets_ul(sec["bullets"])
        elif code:
            ul = bullets_ul(sec["bullets"], style="flex:1.3")
            body = f"""      <div style="display:flex; gap:2em; align-items:flex-start;">
{ul}
        <pre style="flex:1"><code class="language-json">{esc(code)}</code></pre>
      </div>"""
        else:
            body = bullets_ul(sec["bullets"])
        footer = terms_footer_html(sec.get("terms"))
        slides_html.append(f"""    <section>
      <h2>{esc(sec['title'])}</h2>
      <div class="slide-body">
{body}
      </div>{footer}{notes_html}
    </section>""")

    closing = data.get("closing")
    if closing:
        body = bullets_ul(closing["bullets"])
        footer = terms_footer_html(closing.get("terms"))
        slides_html.append(f"""    <section>
      <h2>{esc(closing['title'])}</h2>
      <div class="slide-body">
{body}
      </div>{footer}
    </section>""")

    html_out = HTML_TEMPLATE.format(
        topic=esc(data["topic"]),
        cdn=REVEAL_CDN,
        slides="\n".join(slides_html),
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html_out, encoding="utf-8")


# ---------------------------------------------------------------------------
# Gallery index update
# ---------------------------------------------------------------------------

GALLERY_TEMPLATE = """<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>PT Agent 발표자료 모음</title>
<style>
  :root {{ --accent:#4f8cff; --bg:#0b1020; --card:#141b31; --text:#e8ecf7; --muted:#9aa4c2; }}
  * {{ box-sizing: border-box; }}
  body {{ margin:0; background:var(--bg); color:var(--text); font-family:'Segoe UI',sans-serif; padding:40px 20px; }}
  h1 {{ text-align:center; margin-bottom:8px; }}
  p.sub {{ text-align:center; color:var(--muted); margin-top:0; }}
  .grid {{ display:grid; grid-template-columns:repeat(auto-fill,minmax(280px,1fr)); gap:20px; max-width:1000px; margin:40px auto 0; }}
  .card {{ background:var(--card); border-radius:12px; padding:24px; border:1px solid #262f4d; }}
  .card h2 {{ margin:0 0 8px; font-size:1.15em; color:var(--accent); }}
  .card p {{ color:var(--muted); font-size:0.9em; margin:0 0 16px; }}
  .card a {{ display:inline-block; margin-right:10px; padding:8px 14px; border-radius:6px; text-decoration:none; font-size:0.85em; }}
  .card a.view {{ background:var(--accent); color:#fff; }}
  .card a.dl {{ border:1px solid var(--accent); color:var(--accent); }}
</style>
</head>
<body>
<h1>PT Agent 발표자료 모음</h1>
<p class="sub">주제와 목차만 입력하면 리서치 → PPT → 웹 슬라이드까지 자동 생성됩니다</p>
<div class="grid">
{cards}
</div>
</body>
</html>
"""

CARD_TEMPLATE = """  <div class="card">
    <h2>{topic}</h2>
    <p>{subtitle}{audience}</p>
    <a class="view" href="{slug}/">웹으로 보기</a>
    <a class="dl" href="{slug}/slides.pptx" download>PPT 다운로드</a>
  </div>"""


def update_gallery(entries):
    cards = "\n".join(
        CARD_TEMPLATE.format(
            topic=esc(e["topic"]),
            subtitle=esc(e.get("subtitle", "")),
            audience=f' · 대상: {esc(e["audience"])}' if e.get("audience") else "",
            slug=e["slug"],
        )
        for e in entries
    )
    out = GALLERY_TEMPLATE.format(cards=cards)
    (DOCS_DIR / "index.html").write_text(out, encoding="utf-8")


def discover_existing_entries():
    entries = []
    content_dir = Path(__file__).resolve().parent / "content"
    for p in sorted(content_dir.glob("*.json")):
        d = load_content(p)
        entries.append(
            {
                "slug": d["slug"],
                "topic": d["topic"],
                "subtitle": d.get("subtitle", ""),
                "audience": d.get("audience", ""),
            }
        )
    return entries


def main():
    if len(sys.argv) != 2:
        print("Usage: python build_deck.py <content.json>")
        sys.exit(1)

    content_path = Path(sys.argv[1])
    data = load_content(content_path)
    slug = data["slug"]

    out_dir = DOCS_DIR / slug
    build_pptx(data, out_dir / "slides.pptx")
    build_html(data, out_dir / "index.html")

    update_gallery(discover_existing_entries())

    print(f"Built: {out_dir / 'slides.pptx'}")
    print(f"Built: {out_dir / 'index.html'}")
    print(f"Updated: {DOCS_DIR / 'index.html'}")


if __name__ == "__main__":
    main()
